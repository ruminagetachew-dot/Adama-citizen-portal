"""
Convert Practical Attachment Report from Markdown to DOCX
"""

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
import re

def add_title(doc, text):
    """Add centered title"""
    title = doc.add_heading(text, level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_format = title.runs[0]
    title_format.font.size = Pt(16)
    title_format.font.bold = True
    return title

def add_heading1(doc, text):
    """Add level 1 heading"""
    heading = doc.add_heading(text, level=1)
    heading_format = heading.runs[0]
    heading_format.font.size = Pt(14)
    heading_format.font.bold = True
    heading_format.font.color.rgb = RGBColor(0, 51, 102)
    return heading

def add_heading2(doc, text):
    """Add level 2 heading"""
    heading = doc.add_heading(text, level=2)
    heading_format = heading.runs[0]
    heading_format.font.size = Pt(12)
    heading_format.font.bold = True
    return heading

def add_heading3(doc, text):
    """Add level 3 heading"""
    heading = doc.add_heading(text, level=3)
    heading_format = heading.runs[0]
    heading_format.font.size = Pt(11)
    heading_format.font.bold = True
    return heading

def add_paragraph(doc, text, bold=False, italic=False):
    """Add paragraph with optional formatting"""
    para = doc.add_paragraph()
    run = para.add_run(text)
    run.font.size = Pt(11)
    run.font.name = 'Times New Roman'
    if bold:
        run.font.bold = True
    if italic:
        run.font.italic = True
    return para

def add_table_from_markdown(doc, lines):
    """Convert markdown table to Word table"""
    rows = []
    for line in lines:
        if '|' in line and not line.strip().startswith('|---'):
            cells = [cell.strip() for cell in line.split('|')[1:-1]]
            if cells:
                rows.append(cells)
    
    if rows:
        table = doc.add_table(rows=len(rows), cols=len(rows[0]))
        table.style = 'Light Grid Accent 1'
        
        for i, row_data in enumerate(rows):
            row = table.rows[i]
            for j, cell_data in enumerate(row_data):
                cell = row.cells[j]
                cell.text = cell_data
                # Bold first row (header)
                if i == 0:
                    for paragraph in cell.paragraphs:
                        for run in paragraph.runs:
                            run.font.bold = True
        
        return table
    return None

def process_markdown_to_docx(md_file, docx_file):
    """Main conversion function"""
    doc = Document()
    
    # Set default font
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(11)
    
    # Set margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1.25)
        section.right_margin = Inches(1.25)
    
    # Read markdown file
    with open(md_file, 'r', encoding='utf-8') as f:
        lines = f.readlines()
    
    i = 0
    in_table = False
    table_lines = []
    
    while i < len(lines):
        line = lines[i].rstrip()
        
        # Skip empty lines at start
        if not line and i == 0:
            i += 1
            continue
        
        # Handle tables
        if '|' in line and not in_table:
            in_table = True
            table_lines = [line]
            i += 1
            continue
        
        if in_table:
            if '|' in line:
                table_lines.append(line)
                i += 1
                continue
            else:
                # End of table
                add_table_from_markdown(doc, table_lines)
                doc.add_paragraph()
                in_table = False
                table_lines = []
        
        # Handle headings
        if line.startswith('# '):
            add_title(doc, line[2:].strip())
        elif line.startswith('## '):
            heading_text = line[3:].strip()
            if 'CHAPTER' in heading_text.upper():
                add_heading1(doc, heading_text)
            else:
                add_heading2(doc, heading_text)
        elif line.startswith('### '):
            add_heading2(doc, line[4:].strip())
        elif line.startswith('#### '):
            add_heading3(doc, line[5:].strip())
        
        # Handle bold text
        elif line.startswith('**') and line.endswith('**'):
            add_paragraph(doc, line.strip('*'), bold=True)
        
        # Handle horizontal rules
        elif line.startswith('---'):
            doc.add_paragraph('_' * 80)
        
        # Handle bullet points
        elif line.startswith('- ') or line.startswith('* '):
            para = doc.add_paragraph(line[2:].strip(), style='List Bullet')
            para.paragraph_format.left_indent = Inches(0.5)
        
        # Handle numbered lists
        elif re.match(r'^\d+\.', line):
            text = re.sub(r'^\d+\.\s*', '', line)
            para = doc.add_paragraph(text, style='List Number')
            para.paragraph_format.left_indent = Inches(0.5)
        
        # Handle regular paragraphs
        elif line.strip():
            # Check for bold/italic inline formatting
            clean_line = line
            # Remove markdown bold
            clean_line = re.sub(r'\*\*(.*?)\*\*', r'\1', clean_line)
            # Remove markdown italic
            clean_line = re.sub(r'\*(.*?)\*', r'\1', clean_line)
            # Remove markdown links [text](url)
            clean_line = re.sub(r'\[(.*?)\]\(.*?\)', r'\1', clean_line)
            
            add_paragraph(doc, clean_line)
        
        # Empty line - add spacing
        else:
            if i > 0 and lines[i-1].strip():  # Only add space after content
                doc.add_paragraph()
        
        i += 1
    
    # Save document
    doc.save(docx_file)
    print(f"✅ Successfully converted to: {docx_file}")

if __name__ == "__main__":
    md_file = "c:/Users/Hp/Desktop/Citizen/PRACTICAL_ATTACHMENT_REPORT.md"
    docx_file = "c:/Users/Hp/Desktop/Citizen/PRACTICAL_ATTACHMENT_REPORT.docx"
    
    try:
        process_markdown_to_docx(md_file, docx_file)
        print(f"\n📄 Document created successfully!")
        print(f"📍 Location: {docx_file}")
    except Exception as e:
        print(f"❌ Error: {e}")
        import traceback
        traceback.print_exc()
